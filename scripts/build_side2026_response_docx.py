#!/usr/bin/env python3
"""Render the point-by-point response to reviewers as Response_237.docx.

CMT wants the response as a Word file. The Markdown in the submission folder
is the source of record; this converts the small subset it uses (three heading
levels, block quotes for the reviewer's words, bullet lists, bold runs) so the
two never drift.

Usage:
    python scripts/build_side2026_response_docx.py
    # -> docs/dissertation/side2026_submission/revision/Response_237.docx
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

REPO = Path(__file__).resolve().parents[1]
SUB = REPO / "docs" / "dissertation" / "side2026_submission"
SRC = SUB / "response_to_reviewers.md"
OUT = SUB / "revision" / "Response_237.docx"
FONT = "Times New Roman"


def add_runs(paragraph, text: str, *, italic: bool = False) -> None:
    for i, part in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not part:
            continue
        bold = part.startswith("**")
        run = paragraph.add_run(part.strip("*") if bold else part)
        run.bold = bold
        run.italic = italic
        run.font.name = FONT
        run.font.size = Pt(11)


def build() -> Path:
    doc = Document()
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(11)

    for line in SRC.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            add_runs(p, line[2:], italic=True)
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            add_runs(doc.add_paragraph(), line)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print("written", build())
